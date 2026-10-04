"""PPTX -> segments per slide shape / table row / notes. Deterministic (python-pptx)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..schema import SegmentBuilder, Segment
from ..sources import SourceInfo
from ._common import Ctx, clean


def parse(info: SourceInfo, path: Path, ctx: Ctx) -> list[Segment]:
    b = SegmentBuilder(info.doc_id)
    prs = Presentation(path)
    for sno, slide in enumerate(prs.slides, start=1):
        text_idx = 0
        for shape in slide.shapes:
            loc = {"slide": sno, "shape": shape.name}
            if shape.has_text_frame and clean(shape.text_frame.text):
                text = "\n".join(clean(p.text) for p in shape.text_frame.paragraphs if clean(p.text))
                b.add("slide_text", text, {**loc, "order": text_idx},
                      data={"role": "title" if text_idx == 0 else "body"})
                text_idx += 1
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[clean(c.text) for c in r.cells] for r in shape.table.rows]
                header = rows[0]
                for ri, row in enumerate(rows[1:], start=1):
                    cells = {header[i]: v for i, v in enumerate(row)}
                    b.add("table_row", " | ".join(f"{k}: {v}" for k, v in cells.items()),
                          {**loc, "row": ri}, data={"header": header, "cells": cells, "lists": {}})
            elif shape.shape_type == 13:     # picture: not read here; flagged so the loss is visible
                b.add("metadata", f"picture shape '{shape.name}' not text-extracted", loc,
                      data={"role": "unread_picture"})
        if slide.has_notes_slide and clean(slide.notes_slide.notes_text_frame.text):
            b.add("slide_text", clean(slide.notes_slide.notes_text_frame.text),
                  {"slide": sno, "shape": "notes"}, data={"role": "speaker_notes"})
    return b.segments
